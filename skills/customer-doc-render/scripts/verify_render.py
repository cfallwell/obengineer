#!/usr/bin/env python3
"""Prove a rendered .docx matches its Markdown source and the approved format.

    python3 verify_render.py GUIDE.md GUIDE.docx

Checks, all fence-aware so comments inside code blocks are never counted as
headings:

  structure   heading / table / code-block counts agree with the Markdown
  levels      a top-level Markdown section is a Word Heading 1
  paging      every section heading carries w:pageBreakBefore
  title page  a simple title page, then the contents, and nothing else
  fidelity    no literal Markdown markers leaked into prose runs
  secrets     no credential-shaped value reached the document
  toc         a Table of Contents field exists and Word will refresh it

Exit code 0 when every check passes, 1 otherwise.
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover - dependency guard
    sys.exit("python-docx is required: pip install python-docx")

# Values that must never reach a customer document. Matches the guide's deny list.
SECRET_PATTERNS = [
    ("jwt", re.compile(r"\beyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]+")),
    ("opaque token", re.compile(r"\b[0-9a-f]{32,}\b")),
    ("aws access key", re.compile(r"\b(?:AKIA|ASIA|AGPA|AIDA|AROA)[A-Z0-9]{12,}\b")),
    ("google api key", re.compile(r"\bAIza[0-9A-Za-z_-]{35}\b")),
    ("github token", re.compile(r"\bgh[posur]_[A-Za-z0-9]{20,}\b")),
    ("stripe key", re.compile(r"\b(?:sk|pk)_(?:live|test)_[A-Za-z0-9]{16,}\b")),
    ("private key block", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    ("basic-auth url", re.compile(r"\b[a-z][a-z0-9+.-]*://[^/\s:@]+:[^/\s@]+@")),
]

TOC_HEADING = "Table of Contents"
# A simple title page: title, subtitle, tagline, author, audience, date, version.
MAX_TITLE_PAGE_LINES = 7


def has_page_break_before(paragraph) -> bool:
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is None:
        return False
    el = ppr.find(qn("w:pageBreakBefore"))
    if el is None:
        return False
    return el.get(qn("w:val")) not in ("false", "0", "off")


def parse_markdown(path: Path, section_level: int) -> dict:
    """Fence-aware structural census of the Markdown source."""
    headings: list[tuple[int, str]] = []
    tables = code_blocks = 0
    in_fence = in_table = False

    for line in path.read_text().split("\n"):
        if line.startswith("```"):
            in_fence = not in_fence
            if in_fence:
                code_blocks += 1
            continue
        if in_fence:
            continue
        if line.startswith("|"):
            if not in_table:
                tables += 1
                in_table = True
            continue
        in_table = False
        m = re.match(r"^(#{1,6}) +(.*)$", line)
        if m:
            headings.append((len(m.group(1)), m.group(2)))

    # The leading H1 becomes the title page, so it is not a body heading and
    # does not get a page break of its own.
    body = headings[1:] if headings and headings[0][0] == 1 else list(headings)
    return {
        "headings": headings,
        "body_headings": body,
        "sections": [h for h in body if h[0] <= section_level],
        "tables": tables,
        "code_blocks": code_blocks,
    }


def check(md: Path, docx: Path, section_level: int) -> list[str]:
    failures: list[str] = []
    src = parse_markdown(md, section_level)
    doc = Document(docx)

    # python-docx builds a fresh proxy on every access, so materialise once and
    # compare the underlying XML elements rather than the wrappers.
    paragraphs = list(doc.paragraphs)
    rendered = [p for p in paragraphs if p.style.name.startswith("Heading")]
    toc_headings = [p for p in rendered if p.text.strip() == TOC_HEADING]
    toc_elements = {id(p._p) for p in toc_headings}
    body_rendered = [p for p in rendered if id(p._p) not in toc_elements]

    # ---- structure -------------------------------------------------------
    expected = len(src["body_headings"])
    if len(body_rendered) != expected:
        failures.append(
            f"heading count: markdown has {expected} body headings, "
            f"docx has {len(body_rendered)}"
        )

    grid = [t for t in doc.tables if t.style and t.style.name == "Table Grid"]
    code_boxes = len(doc.tables) - len(grid)
    if len(grid) != src["tables"]:
        failures.append(
            f"table count: markdown has {src['tables']}, docx has {len(grid)}"
        )
    if code_boxes != src["code_blocks"]:
        failures.append(
            f"code-block count: markdown has {src['code_blocks']}, "
            f"docx has {code_boxes}"
        )

    # ---- heading levels --------------------------------------------------
    # A top-level Markdown section must land on Word Heading 1, or the contents
    # list and the navigation pane are indented a level too deep.
    if section_level >= 2 and len(body_rendered) == expected:
        wrong = [
            f"{text[:40]!r} -> {para.style.name}"
            for (level, text), para in zip(src["body_headings"], body_rendered)
            if para.style.name != f"Heading {max(1, min(level - 1, 4))}"
        ]
        if wrong:
            failures.append(f"heading level mapping: {wrong[:3]}")

    # ---- paging ----------------------------------------------------------
    sections = [p for p in body_rendered if p.style.name == "Heading 1"]
    if len(sections) != len(src["sections"]):
        failures.append(
            f"section count: {len(src['sections'])} in markdown, "
            f"{len(sections)} Heading 1 paragraphs in docx"
        )
    unbroken = [p.text[:60] for p in sections if not has_page_break_before(p)]
    if unbroken:
        failures.append(
            f"section headings without w:pageBreakBefore: {unbroken[:3]}")
    if toc_headings and not has_page_break_before(toc_headings[0]):
        failures.append("the contents page does not start on a new page")

    # ---- title page ------------------------------------------------------
    if toc_headings:
        toc_index = next(i for i, p in enumerate(paragraphs)
                         if id(p._p) == id(toc_headings[0]._p))
        title_lines = [p.text.strip() for p in paragraphs[:toc_index]
                       if p.text.strip()]
        if len(title_lines) > MAX_TITLE_PAGE_LINES:
            failures.append(
                f"title page has {len(title_lines)} lines of text; the approved "
                f"format allows at most {MAX_TITLE_PAGE_LINES}. Move prose into "
                "a body section."
            )
        long_lines = [t[:60] for t in title_lines if len(t) > 160]
        if long_lines:
            failures.append(f"prose paragraph on the title page: {long_lines[:2]}")

    # ---- fidelity: markers that leaked as literal text -------------------
    def prose_paragraphs():
        for p in doc.paragraphs:
            yield p
        for t in grid:  # code boxes legitimately contain backticks and asterisks
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    leaked_bold = [p.text[:70] for p in prose_paragraphs() if "**" in p.text]
    if leaked_bold:
        failures.append(f"literal '**' in rendered prose: {leaked_bold[:3]}")

    leaked_code = [
        p.text[:70]
        for p in prose_paragraphs()
        if "`" in p.text
        and not all(r.font.name == "Consolas" for r in p.runs if r.text.strip())
    ]
    if leaked_code:
        failures.append(f"literal backticks in rendered prose: {leaked_code[:3]}")

    # ---- secrets ---------------------------------------------------------
    everything = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                everything += "\n" + cell.text
    for label, pattern in SECRET_PATTERNS:
        hits = pattern.findall(everything)
        if hits:
            failures.append(f"possible {label} in document ({len(hits)} match(es))")

    # ---- toc -------------------------------------------------------------
    with zipfile.ZipFile(docx) as z:
        document_xml = z.read("word/document.xml").decode("utf-8", "replace")
        settings_xml = z.read("word/settings.xml").decode("utf-8", "replace")
    if "TOC" not in document_xml:
        failures.append("no Table of Contents field found")
    if "updateFields" not in settings_xml:
        failures.append("updateFields not set — Word will not refresh the contents")

    return failures


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("docx", type=Path)
    ap.add_argument("--section-level", type=int, default=2, choices=(1, 2))
    args = ap.parse_args(argv)

    for p in (args.markdown, args.docx):
        if not p.is_file():
            sys.exit(f"not a file: {p}")

    failures = check(args.markdown, args.docx, args.section_level)
    if failures:
        print(f"FAIL — {len(failures)} problem(s):")
        for f in failures:
            print(f"  - {f}")
        return 1
    print(f"PASS — {args.docx.name} matches {args.markdown.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
